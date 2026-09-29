from pathlib import Path

import pdfplumber
import pymupdf
from docx import Document
from pypdf import PdfReader
from rapidocr_onnxruntime import RapidOCR
from zipfile import ZipFile
from xml.etree import ElementTree

def extract_text_from_pdf(file_path):
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="ignore")

    if suffix == ".docx":
        document = Document(str(path))
        parts = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]
        for section in document.sections:
            parts.extend(p.text for p in section.header.paragraphs if p.text.strip())
            parts.extend(p.text for p in section.footer.paragraphs if p.text.strip())

        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))

        if parts:
            return "\n".join(parts)

        # Fallback for text stored in other parts of the Word file.
        with ZipFile(path) as docx_file:
            for name in docx_file.namelist():
                if name.startswith("word/") and name.endswith(".xml"):
                    root = ElementTree.fromstring(docx_file.read(name))
                    parts.extend(
                        item.text
                        for item in root.iter()
                        if item.tag.endswith("}t") and item.text
                    )

        return "\n".join(parts)

    if suffix != ".pdf":
        raise ValueError("Only PDF, TXT, and DOCX files are supported.")

    # First try extracting selectable text with pypdf.
    try:
        reader = PdfReader(str(path))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)

        if text.strip():
            return text
    except Exception:
        pass

    # Try pdfplumber if pypdf did not find text.
    try:
        with pdfplumber.open(str(path)) as pdf:
            text = "\n".join(
                page_text
                for page in pdf.pages
                if (page_text := page.extract_text())
            )

        if text.strip():
            return text
    except Exception:
        pass

    # Use OCR for scanned or image-based PDFs.
    try:
        ocr = RapidOCR()
        pages_text = []

        with pymupdf.open(str(path)) as doc:
            for page_num, page in enumerate(doc):
                pix = page.get_pixmap(dpi=220)
                img_path = path.with_name(
                    f"{path.stem}_page_{page_num + 1}.png"
                )

                try:
                    pix.save(str(img_path))
                    result, _ = ocr(str(img_path))

                    if result:
                        pages_text.extend(item[1] for item in result)
                finally:
                    img_path.unlink(missing_ok=True)

        text = "\n".join(pages_text)
        if text.strip():
            return text
    except Exception:
        pass

    raise ValueError(
        "No readable text could be extracted from this file. "
        "Please upload a file that contains selectable text or a readable scan."
    )